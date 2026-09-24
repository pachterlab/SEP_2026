#!/usr/bin/env python3
"""Check Python/R dependencies and real LaTeX figure rendering."""

import importlib
from pathlib import Path
import shutil
import sys
import subprocess
from tempfile import TemporaryDirectory


def check_latex():
    """Verify the external TeX tools and fonts used by the methods figure."""
    # Use the same user-local dvipng fallback as the figure script, if needed.
    import fig1_architecture
    for name in ("latex", "dvipng", "kpsewhich"):
        executable = shutil.which(name)
        if executable is None:
            raise RuntimeError(
                f"Missing {name}. On Debian/Ubuntu run make system-deps; "
                "on other systems install the TeX dependencies listed in README.md."
            )
        print(f"{name}: {executable}", flush=True)
    for filename in ("type1cm.sty", "type1ec.sty"):
        found = subprocess.run(["kpsewhich", filename], capture_output=True, text=True)
        if found.returncode or not found.stdout.strip():
            raise RuntimeError(
                f"Missing {filename}. Install the complete LaTeX dependencies "
                "with make system-deps (Debian/Ubuntu); see README.md."
            )
        print(f"{filename}: {found.stdout.strip()}", flush=True)


def main():
    import matplotlib
    matplotlib.use("Agg")
    check_latex()
    from lib.figure_style import figure_font, font_rcparams
    matplotlib.rcParams.update(font_rcparams())
    print(f"Figure font: {figure_font()}", flush=True)
    for name in ("numpy", "pandas", "scipy", "sklearn", "matplotlib", "PIL",
                 "openpyxl", "umap", "elpigraph", "clrcycle"):
        module = importlib.import_module(name)
        print(f"{name}: {getattr(module, '__version__', 'OK')}", flush=True)
    prefix = Path(sys.prefix).resolve()
    rscript = shutil.which("Rscript")
    if not rscript or not Path(rscript).resolve().is_relative_to(prefix):
        raise RuntimeError("Activate the complete project environment; Rscript must come from it. Run make environment.")
    code = r'''
    prefix <- normalizePath(commandArgs(TRUE)[1])
    for (pkg in c("Matrix", "Biobase", "CHIRAL")) {
        suppressPackageStartupMessages(library(pkg, character.only=TRUE))
        stopifnot(startsWith(normalizePath(find.package(pkg)), paste0(prefix, "/")))
        cat(pkg, as.character(packageVersion(pkg)), "\n")
    }
    set.seed(17)
    phases <- seq(0, 2*pi, length.out=25)[-25]
    mat <- 10 + outer(phases, phases, function(a,b) cos(a-b))
    rownames(mat) <- paste0("gene", seq_len(nrow(mat)))
    colnames(mat) <- paste0("sample", seq_len(ncol(mat)))
    out <- CHIRAL::CHIRAL(mat, clockgenes=rownames(mat), pbar=FALSE)
    stopifnot(length(out$phi) == ncol(mat), all(is.finite(out$phi)))
    stopifnot(identical(Biobase::exprs(Biobase::ExpressionSet(mat)), mat))
    stopifnot(all(as.matrix(Matrix::Matrix(mat, sparse=TRUE)) == mat))
    cat(R.version.string, "- R checks passed\n")
    '''
    from lib.common import run_r
    run_r(code, prefix, executable=rscript)
    from openpyxl import Workbook, load_workbook
    import matplotlib.pyplot as plt
    with TemporaryDirectory() as folder:
        path = Path(folder)
        book = Workbook()
        book.active.append(["dependency check", 1])
        book.save(path / "check.xlsx")
        loaded = load_workbook(path / "check.xlsx")
        assert loaded.active["B1"].value == 1
        loaded.close()
        fig, ax = plt.subplots()
        # Exercise the actual methods-panel formulas and the external TeX
        # renderer, rather than merely exporting a figure containing no math.
        from fig1_architecture import draw_projection, workflow_values
        values = workflow_values()
        draw_projection(ax, values["display_order"], values["projected"])
        for suffix in ("pdf", "svg", "png"):
            fig.savefig(path / f"check.{suffix}")
        plt.close(fig)
    print("Environment checks passed: R, CHIRAL, Excel, and LaTeX PDF/SVG/PNG export.")


if __name__ == "__main__":
    main()
