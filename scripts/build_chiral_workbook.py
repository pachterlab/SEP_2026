"""Package the saved CHIRAL comparison as a standalone supplementary workbook."""

import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import CellIsRule

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "results/chiral_comparison"
WORKBOOK = ROOT / "results/supplementary_data/Supplementary_Data_4_CHIRAL_vs_clrcycle.xlsx"
NAVY = "17365D"
TEAL = "DDEEF0"


def format_table(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 46
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[cell.column_letter].width = 23
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.000"
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F0F5FA")


def write_chiral_workbook(path=WORKBOOK):
    """Write a summary first, followed by tissue results, phases, and methods."""
    table = pd.read_csv(SOURCE / "chiral_vs_clrcycle.csv", float_precision="round_trip")
    phases = pd.read_csv(SOURCE / "sample_phases.csv", float_precision="round_trip")
    audit = json.loads((SOURCE / "comparison_audit.json").read_text())
    summary = pd.DataFrame([
        {
            "method": "CHIRAL" if method == "chiral" else method,
            "mean_R": table[f"{method}_R"].mean(),
            "mean_abs_err_h": table[f"{method}_mean_abs_err_h"].mean(),
            "tissues_with_higher_R": int((table[f"{method}_R"] > table[f"{other}_R"] + 1e-12).sum()),
            "tissues_with_lower_error": int((table[f"{method}_mean_abs_err_h"] < table[f"{other}_mean_abs_err_h"] - 1e-12).sum()),
        }
        for method, other in (("chiral", "clrcycle"), ("clrcycle", "chiral"))
    ])
    pd.testing.assert_frame_equal(summary, pd.read_csv(SOURCE / "summary.csv"))
    methods = pd.DataFrame([
        ("Dataset", "GSE54650; current Figure 1 circadian tissue panels."),
        ("Panel selection", audit["selection"]),
        ("Preprocessing", audit["preprocessing"]),
        ("Alignment", audit["alignment"]),
        ("R (higher is better)", audit["R"] + "; phase agreement, not a rank-only ordering score."),
        ("Phase error (lower is better)", audit["error"]),
        ("Overall means", audit["aggregation"]),
        ("R difference", "clrcycle minus CHIRAL; positive favors clrcycle."),
        ("Error difference", "clrcycle minus CHIRAL, in hours; negative favors clrcycle."),
        ("Scope", audit["scope"]),
        ("Panel source", audit["panel_source"]),
        ("Coordinate source", audit["coordinate_source"]),
        ("Saved results", "results/chiral_comparison/: chiral_vs_clrcycle.csv, summary.csv, sample_phases.csv, comparison_audit.json"),
        ("Display precision", "Cells display three decimal places; underlying numeric values retain source precision."),
    ], columns=["Item", "Explanation"])
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        book = writer.book
        overview = book.create_sheet("Summary")
        table.to_excel(writer, sheet_name="Tissue_comparison", index=False)
        phases.to_excel(writer, sheet_name="Sample_phases", index=False)
        methods.to_excel(writer, sheet_name="Methods", index=False)
        for sheet in book.worksheets[1:]:
            format_table(sheet)
        overview.sheet_view.showGridLines = False
        overview.merge_cells("A1:E1")
        overview["A1"] = "CHIRAL versus clrcycle"
        overview["A1"].font = Font(size=20, bold=True, color=NAVY)
        overview.row_dimensions[1].height = 32
        overview.merge_cells("A2:E2")
        overview["A2"] = (
            f"GSE54650 | {len(table)} tissues | {int(table.n_probes.min())} probes per tissue | "
            f"{int(table.n_samples.sum())} samples total | matched inputs and scoring"
        )
        overview["A2"].alignment = Alignment(wrap_text=True)
        overview.row_dimensions[2].height = 32
        headers = ["Method", "Mean R\nHigher is better", "Mean phase error (h)\nLower is better",
                   "Tissues with higher R", "Tissues with lower error"]
        for col, label in enumerate(headers, 1):
            cell = overview.cell(4, col, label)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            overview.column_dimensions[cell.column_letter].width = 25
        overview.row_dimensions[4].height = 36
        for row, values in enumerate(summary.itertuples(index=False, name=None), 5):
            for col, value in enumerate(values, 1):
                cell = overview.cell(row, col, value)
                cell.fill = PatternFill("solid", fgColor=TEAL)
                cell.font = Font(size=12, bold=col == 1)
                if col in (2, 3):
                    cell.number_format = "0.000"
            overview.row_dimensions[row].height = 25
        better = table.loc[
            (table.clrcycle_R > table.chiral_R + 1e-12)
            & (table.clrcycle_mean_abs_err_h < table.chiral_mean_abs_err_h - 1e-12), "tissue"
        ].tolist()
        notes = [
            (8, "How to read this comparison"),
            (9, "Overall values are equally weighted means across tissues. R measures phase agreement; error measures placement accuracy in hours."),
            (11, "clrcycle has higher R and lower phase error in: " + (", ".join(better) or "none") + "."),
            (13, "This comparison uses the same selected probes and samples for both methods. A single CHIRAL initialization was evaluated; these results do not establish general robustness."),
            (16, "Tissue_comparison: all tissue-level scores and paired differences"),
            (17, "Sample_phases: raw phases, aligned phases, and sample-level errors"),
            (18, "Methods: metric definitions, panel selection, and interpretation"),
        ]
        for row, value in notes:
            overview.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            cell = overview.cell(row, 1, value)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.font = Font(size=11, bold=row == 8, color=NAVY)
            overview.row_dimensions[row].height = 38 if row in (9, 11, 13) else 24
        for row, target in ((16, "Tissue_comparison"), (17, "Sample_phases"), (18, "Methods")):
            overview.cell(row, 1).hyperlink = f"#'{target}'!A1"
        overview.freeze_panes = "A5"
        overview.sheet_properties.pageSetUpPr.fitToPage = True
        overview.page_setup.orientation = "landscape"
        overview.page_setup.paperSize = overview.PAPERSIZE_A4
        overview.page_setup.fitToWidth = 1
        overview.page_setup.fitToHeight = 1
        overview.print_area = "A1:E18"
        tissue_sheet = book["Tissue_comparison"]
        for col, operator in (("I", "greaterThan"), ("J", "lessThan")):
            tissue_sheet.conditional_formatting.add(
                f"{col}2:{col}{len(table) + 1}",
                CellIsRule(operator=operator, formula=["0"],
                           fill=PatternFill("solid", fgColor="E2F0D9")),
            )
        book["Methods"].column_dimensions["A"].width = 32
        book["Methods"].column_dimensions["B"].width = 100
        for row in book["Methods"].iter_rows(min_row=2):
            book["Methods"].row_dimensions[row[0].row].height = 48
            row[1].alignment = Alignment(wrap_text=True, vertical="center")
    print(f"Wrote {path}")
    return path


if __name__ == "__main__":
    write_chiral_workbook()
